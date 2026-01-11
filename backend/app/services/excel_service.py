import openpyxl
from typing import Dict, List

class ExcelTemplateService:
    def __init__(self):
        pass

    def load_template(self, file_path: str):
        """
        Carga el libro de Excel manteniendo estilos (data_only=False).
        """
        try:
            # openpyxl es clave aquí porque permite editar celdas sin destruir estilos
            self.workbook = openpyxl.load_workbook(file_path)
            return True
        except Exception as e:
            print(f"Error cargando plantilla: {e}")
            return False

    def get_sheet_names(self) -> List[str]:
        if not self.workbook:
            return []
        return self.workbook.sheetnames

    def write_data(self, sheet_name: str, data_map: Dict[str, str], output_path: str):
        """
        Escribe datos en celdas específicas.
        data_map ejemplo: {'B2': 'Juan Perez', 'C10': '12345'}
        """
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Hoja {sheet_name} no existe")
            
        ws = self.workbook[sheet_name]
        
        for cell_address, value in data_map.items():
            # Aquí es donde ocurre la magia: solo inyectamos el valor
            # El estilo de la celda permanece intacto
            ws[cell_address] = value
            
        self.workbook.save(output_path)
        return output_path
